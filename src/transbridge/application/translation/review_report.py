"""Present advisory AI review evidence without treating it as a commit decision."""

from __future__ import annotations

from collections.abc import Callable, Mapping
import json

from .postprocess import ReportSnapshot

REVIEW_STATUS_LABELS = {
    "revised": "AI 已修订",
    "retained": "保留原译：语境适用",
    "term_suspect": "疑似参考术语错误",
    "needs_review": "待复核",
}
ISSUE_TYPE_LABELS = {
    "term_mismatch": "术语译法差异",
    "name_inconsistency": "同名译法差异",
    "untranslated": "疑似未翻译",
    "missing_translation": "疑似漏译",
    "syntax_error": "标记或语法疑点",
}


def review_label(details: Mapping[str, object]) -> str:
    return REVIEW_STATUS_LABELS.get(str(details.get("review_status", "")), "")


def review_issue_types(details: Mapping[str, object]) -> str:
    issues = details.get("detected_issues") or ()
    names = dict.fromkeys(
        str(issue.get("issue_type") or issue.get("type") or issue.get("code") or "质量疑点")
        for issue in issues
        if isinstance(issue, Mapping)
    )
    return "、".join(ISSUE_TYPE_LABELS.get(name, name) for name in names)


def review_issue_text(details: Mapping[str, object]) -> str:
    issues = details.get("detected_issues") or ()
    return "\n".join(
        json.dumps(issue, ensure_ascii=False, sort_keys=True, allow_nan=False)
        for issue in issues
        if isinstance(issue, Mapping)
    )


def add_review_sheet(workbook, snapshot: ReportSnapshot, safe_value: Callable[[object], object]) -> None:
    """Readable counterpart to the lossless machine-oriented Entries sheet."""
    from openpyxl.styles import PatternFill

    reviewed = tuple(candidate for candidate in snapshot.candidates if review_label(dict(candidate.report_details)))
    if not reviewed:
        return
    sheet = workbook.create_sheet("校对标注")
    sheet.append((
        "来源",
        "词条 ID",
        "原文",
        "校对前译文",
        "当前译文",
        "AI 校对状态",
        "AI 判断理由",
        "疑点类型",
        "检测疑点（参考线索）",
        "结果可应用",
        "上下文",
    ))
    colors = {"needs_review": "FFF2CC", "term_suspect": "FCE4D6", "revised": "E2F0D9", "retained": "EAF1FB"}
    fills = {status: PatternFill("solid", fgColor=color) for status, color in colors.items()}
    for row_index, candidate in enumerate(reviewed, 2):
        details = dict(candidate.report_details)
        key = candidate.entry_key.to_dict()
        sheet.append(
            tuple(
                safe_value(value)
                for value in (
                    key.get("namespace", ""),
                    key.get("local_key", ""),
                    candidate.original,
                    details.get("review_before_text", candidate.before_text),
                    candidate.text,
                    review_label(details),
                    details.get("review_reason", ""),
                    review_issue_types(details),
                    review_issue_text(details),
                    "是" if candidate.accepted else "否",
                    candidate.context,
                )
            )
        )
        for column in range(1, 12):
            sheet.cell(row_index, column).fill = fills[str(details["review_status"])]
    workbook["Summary"].append(("校对标注说明", "AI 判断供参考；疑点和待复核标记不等于翻译失败，也不代表人工确认。"))
    workbook.active = workbook.index(sheet)


def style_report_workbook(workbook) -> None:
    """Bound widths, freeze headings and filter result rows without scanning cell text."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wrapped = Alignment(vertical="top", wrap_text=True)
    heading_alignment = Alignment(vertical="center", wrap_text=True)
    heading_font = Font(bold=True, color="FFFFFF")
    heading_fill = PatternFill("solid", fgColor="315A7D")
    for sheet in workbook:
        if sheet.title == "Summary":
            sheet.column_dimensions["A"].width = 28
            sheet.column_dimensions["B"].width = 80
            for row in sheet:
                for cell in row:
                    cell.alignment = wrapped
            continue
        sheet.freeze_panes = "C2" if sheet.title in {"Entries", "校对标注"} else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = heading_font
            cell.fill = heading_fill
            cell.alignment = heading_alignment
        for index, cell in enumerate(sheet[1], 1):
            label = str(cell.value)
            width = (
                44
                if label
                in {
                    "original",
                    "before",
                    "candidate",
                    "note",
                    "issues",
                    "review_reason",
                    "detected_issues",
                    "原文",
                    "校对前译文",
                    "当前译文",
                    "AI 判断理由",
                    "检测疑点（参考线索）",
                    "上下文",
                    "message",
                }
                else 24
            )
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrapped
