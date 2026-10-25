"""AI review annotations survive report exports without changing commit decisions."""

import csv
from dataclasses import replace
from io import BytesIO, StringIO
import json

from openpyxl import load_workbook

from tests.contracts.translation.test_polish_report import _Entry
from transbridge.application.contracts import OperationOutcome
from transbridge.application.io import EntryKey, EntryRevision, SourceNamespace
from transbridge.application.translation.polish_report import build_polish_report_snapshot
from transbridge.application.translation.postprocess import PostProcessCandidate, ReportSnapshot
from transbridge.application.translation.postprocess_report import (
    CsvReportRenderer,
    ExcelReportRenderer,
    JsonReportRenderer,
)
from transbridge.application.translation.review_report import REVIEW_STATUS_LABELS
from transbridge.application.translation.task_history import restore_report_snapshot
from transbridge.ui.tools.ai_translator.reporting import report_overview


def review_snapshot():
    candidates = tuple(
        PostProcessCandidate(
            run_id="advisory-run",
            entry_key=EntryKey(SourceNamespace("BSHeartland"), f"name-{index}"),
            before_revision=EntryRevision(),
            original="Bruma guards fighting predators",
            before_text="布鲁玛卫兵对抗掠食者",
            text="布鲁玛卫兵对抗掠食者" if status != "revised" else "布鲁玛卫兵与猛兽交战",
            stage=2,
            accepted=True,
            report_details=tuple(
                {
                    "result_status": "accepted",
                    "review_status": status,
                    "review_reason": "结合语境判断；Guards 是卫兵，参考译法不适用",
                    "detected_issues": [{"issue_type": "term_mismatch", "term": "Guards", "translation": "瓜尔"}],
                }.items()
            ),
        )
        for index, status in enumerate(REVIEW_STATUS_LABELS)
    )
    return ReportSnapshot(
        schema="transbridge.postprocess-report.v1",
        run_id="advisory-run",
        outcome=OperationOutcome.COMPLETED,
        input_count=4,
        accepted_count=4,
        candidates=candidates,
        stage_outcomes=(),
        diagnostics=(),
    )


def test_all_formats_preserve_review_evidence_and_keep_pending_distinct_from_failure():
    snapshot = review_snapshot()
    restored = restore_report_snapshot(json.loads(JsonReportRenderer().render(snapshot).content))
    assert restored.failure_count == 0
    assert restored.accepted_count == 4
    assert report_overview(restored).needs_review == 1
    rows = list(csv.DictReader(StringIO(CsvReportRenderer().render(restored).content.decode("utf-8"))))
    assert {row["review_label"] for row in rows} == set(REVIEW_STATUS_LABELS.values())
    assert all(row["accepted"] == "1" and row["result_status"] == "accepted" for row in rows)
    assert all(json.loads(row["detected_issues"])[0]["translation"] == "瓜尔" for row in rows)
    assert all(row["namespace"] == "BSHeartland" and row["original"] for row in rows)
    workbook = load_workbook(BytesIO(ExcelReportRenderer().render(restored).content))
    sheet = workbook["校对标注"]
    assert workbook.active is sheet
    assert sheet.freeze_panes == "C2"
    assert sheet.auto_filter.ref == "A1:K5"
    labels = {row[5].value: row for row in tuple(sheet)[1:]}
    assert set(labels) == set(REVIEW_STATUS_LABELS.values())
    assert labels["待复核"][5].fill.fgColor.rgb == "00FFF2CC"
    assert labels["保留原译：语境适用"][3].value == labels["保留原译：语境适用"][4].value
    assert labels["AI 已修订"][3].value != labels["AI 已修订"][4].value
    assert labels["疑似参考术语错误"][9].value == "是"
    assert sheet["G2"].alignment.wrap_text
    assert sheet.column_dimensions["G"].width == 44


def test_review_annotations_are_formula_safe_and_json_keeps_original_text():
    snapshot = review_snapshot()
    candidate = snapshot.candidates[0]
    details = dict(candidate.report_details)
    details["review_reason"] = '=HYPERLINK("https://invalid")'
    details["review_before_text"] = "+before"
    candidate = replace(candidate, text="@after", report_details=tuple(details.items()))
    snapshot = replace(snapshot, candidates=(candidate,))
    workbook = load_workbook(BytesIO(ExcelReportRenderer().render(snapshot).content))
    sheet = workbook["校对标注"]
    assert sheet["G2"].value.startswith("'=HYPERLINK")
    assert sheet["D2"].value == "'+before"
    assert sheet["E2"].value == "'@after"
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    row = next(csv.DictReader(StringIO(CsvReportRenderer().render(snapshot).content.decode("utf-8"))))
    assert row["review_reason"].startswith("'=HYPERLINK")
    parsed = json.loads(JsonReportRenderer().render(snapshot).content)
    assert parsed["entries"][0]["report_details"]["review_reason"] == details["review_reason"]


def test_standalone_polish_projection_preserves_advisory_result():
    results = {
        "guard": {
            "original_translation": "卫兵",
            "polished_translation": "卫兵",
            "confidence": 0.9,
            "review_status": "term_suspect",
            "review_reason": "参考术语不适用",
            "detected_issues": [{"issue_type": "term_mismatch", "term": "Guards", "translation": "瓜尔"}],
        }
    }
    snapshot = build_polish_report_snapshot(results, [_Entry("guard", "Guards", "卫兵")], run_id="polish-review")
    details = dict(snapshot.candidates[0].report_details)
    assert details["review_status"] == "term_suspect"
    assert details["review_reason"] == "参考术语不适用"
    assert details["detected_issues"] == results["guard"]["detected_issues"]
    assert snapshot.failure_count == 0 and snapshot.accepted_count == 1
