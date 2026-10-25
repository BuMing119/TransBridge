"""Supplemental AI review must publish stage-only changes and preserve their evidence."""

from io import BytesIO
import json
import threading
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from transbridge.ai_translator.post_processor.proofread_pipeline import ProofreadPipeline
from transbridge.ai_translator.translation_entry_outcomes import TranslationEntryOutcome
from transbridge.application.io import EntryKey, SourceNamespace
from transbridge.application.translation.ai_request_budget import AiRequestBudget
from transbridge.application.translation.entry_alias import ai_entry_key
from transbridge.application.translation.polish_report import build_polish_report_snapshot
from transbridge.application.translation.postprocess_report import ExcelReportRenderer
from transbridge.application.translation.proofread_stage import ProofreadStage
from transbridge.config.llm import LLMConfig
from transbridge.converter.translation_entry import STAGE_QUESTIONABLE, STAGE_TRANSLATED, TranslationEntry
from transbridge.converter.translation_entry_collection import TranslationEntryCollection
from transbridge.ui.tools.ai_translator import source_reference_review
from transbridge.ui.tools.ai_translator.source_execution import SourceOutcome
from transbridge.ui.tools.ai_translator.task_entry_results import TaskEntryResults
from transbridge.ui.tools.ai_translator.task_scope import SourceTask
from transbridge.ui.tools.ai_translator.task_term_closure import close_task_terms
from transbridge.ui.tools.ai_translator.workflow_log_store import WorkflowLogStore


def test_unchanged_translation_pending_review_commits_stage_and_exports_mark(tmp_path, monkeypatch):
    entry = TranslationEntry(
        "guards",
        "guards",
        "Guards",
        "卫兵",
        STAGE_TRANSLATED,
        "NPC_:FULL",
        entry_key=EntryKey(SourceNamespace("plugin"), "guards"),
    )
    collection = TranslationEntryCollection((entry,))
    task = SourceTask("plugin", "Plugin", str(tmp_path / "Plugin.esp"), collection, (entry,), ())
    ledger = TaskEntryResults((task,), "stage-review")
    before_text, before_revision = entry.translation, entry.revision.value
    calls, progress = [], []

    class Provider:
        def chat(self, messages, max_tokens=0):
            items = json.loads(messages[-1]["content"])["entries"]
            calls.append(items)
            return json.dumps(
                {
                    "results": [
                        {
                            "entry_key": item["entry_key"],
                            "final_translation": item["current_translation"],
                            "review_status": "needs_review",
                            "review_reason": "参考译法与当前语境存在冲突，保留译文并标注复核。",
                        }
                        for item in items
                    ]
                },
                ensure_ascii=False,
            )

    monkeypatch.setattr(source_reference_review, "create_llm_client", lambda _: Provider())
    monkeypatch.setattr(
        source_reference_review,
        "WorkflowLogStore",
        lambda path, **kwargs: WorkflowLogStore(path, log_base=tmp_path / "logs", **kwargs),
    )
    pause = threading.Event()
    pause.set()
    budget = AiRequestBudget(1)
    executor = SimpleNamespace(
        attempt_id=None,
        request=SimpleNamespace(run_id="stage-review", request_budget=budget),
        config=LLMConfig(model="fake", max_concurrent=1),
        stop=threading.Event(),
        pause=pause,
        log=lambda *_: None,
        progress=lambda *args: progress.append(args),
    )

    results = source_reference_review.review_source_references(
        executor,
        task,
        {ai_entry_key(entry): {"Guards": "瓜尔"}},
        {},
    )

    current = collection.get(entry.identity)
    result = results[ai_entry_key(entry)]
    assert current.translation == result.text == before_text
    assert current.stage == result.stage == STAGE_QUESTIONABLE
    assert current.revision.value == before_revision + 1
    assert result.status == "succeeded"
    assert dict(result.review_details)["review_status"] == "needs_review"
    assert len(calls) == 1
    assert calls[0][0]["detected_issues"][0]["issue_type"] == "term_mismatch"
    assert budget.snapshot().in_flight == 0
    review_progress = [event for event in progress if event[1] == "term_closure"]
    assert review_progress[-1][2:4] == (1, 1)
    assert progress[-1][1:4] == ("term_closure_finalize", 1, 1)

    outcome = SourceOutcome(task, translation=SimpleNamespace(entry_outcomes={entry.identity: result}))
    ledger.ingest((outcome,), (entry.identity,))
    snapshot = ledger.snapshot(task.key)
    assert snapshot.failure_count == 0 and snapshot.accepted_count == 1
    workbook = load_workbook(BytesIO(ExcelReportRenderer().render(snapshot).content))
    row = tuple(workbook["校对标注"].values)[1]
    assert row[3] == row[4] == before_text
    assert row[5] == "待复核" and row[9] == "是"


def test_supplemental_review_reports_each_completed_batch_before_draft_commit(tmp_path, monkeypatch):
    entries = tuple(
        TranslationEntry(
            str(index),
            str(index),
            "Guard",
            "卫兵",
            STAGE_TRANSLATED,
            "NPC_:FULL",
            entry_key=EntryKey(SourceNamespace("plugin"), str(index)),
        )
        for index in range(8)
    )
    task = SourceTask(
        "plugin", "Plugin", str(tmp_path / "Plugin.esp"), TranslationEntryCollection(entries), entries, ()
    )
    progress = []

    class Provider:
        def chat(self, messages, max_tokens=0):
            items = json.loads(messages[-1]["content"])["entries"]
            return json.dumps(
                {
                    "results": [
                        {
                            "entry_key": item["entry_key"],
                            "final_translation": item["current_translation"],
                            "review_status": "retained",
                            "review_reason": "当前指守卫角色，保留译文。",
                        }
                        for item in items
                    ]
                },
                ensure_ascii=False,
            )

    monkeypatch.setattr(source_reference_review, "create_llm_client", lambda _: Provider())
    monkeypatch.setattr(
        source_reference_review, "ProofreadStage", lambda *args, **kwargs: ProofreadStage(*args, max_items=2, **kwargs)
    )
    monkeypatch.setattr(
        source_reference_review,
        "WorkflowLogStore",
        lambda path, **kwargs: WorkflowLogStore(path, log_base=tmp_path / "logs", **kwargs),
    )
    pause = threading.Event()
    pause.set()
    executor = SimpleNamespace(
        attempt_id=None,
        request=SimpleNamespace(run_id="batch-progress", request_budget=AiRequestBudget(2)),
        config=LLMConfig(model="fake", max_concurrent=2),
        stop=threading.Event(),
        pause=pause,
        log=lambda *_: None,
        progress=lambda *args: progress.append(args),
    )
    results = source_reference_review.review_source_references(
        executor,
        task,
        {ai_entry_key(entry): {"Guard": "瓜尔"} for entry in entries},
        {},
    )
    counts = [event[2] for event in progress if event[1] == "term_closure"]
    assert any(0 < count < len(entries) for count in counts)
    assert counts[-1] == len(entries)
    assert progress[-1][1:4] == ("term_closure_finalize", 8, 8)
    assert len(results) == len(entries)
    assert all(item.status == "succeeded" and item.text == "卫兵" for item in results.values())


def test_latest_retained_review_uses_its_input_after_prior_polish_change(tmp_path):
    original = TranslationEntry(
        "guards",
        "guards",
        "Guards",
        "守备",
        STAGE_TRANSLATED,
        "NPC_:FULL",
        entry_key=EntryKey(SourceNamespace("plugin"), "guards"),
    )
    current = TranslationEntry(
        original.id,
        original.key,
        original.original,
        "卫兵",
        original.stage,
        original.context,
        entry_key=original.identity,
    )
    collection = TranslationEntryCollection((current,))
    task = SourceTask("plugin", "Plugin", str(tmp_path / "Plugin.esp"), collection, (), (original,))
    earlier = ProofreadPipeline._proofread_result(
        original,
        valid=True,
        translation=current.translation,
        review_details={
            "review_status": "revised",
            "review_reason": "按人物语境修正。",
            "review_contract_version": 1,
            "review_before_text": original.translation,
        },
    )
    assert earlier.review_before_text == original.translation
    outcome = SourceOutcome(task, polish={original.id: earlier})
    seen = []

    def supplement(selected_task, references):
        seen.append(references)
        selected = selected_task.collection.get(original.identity)
        return {
            ai_entry_key(original): TranslationEntryOutcome(
                original.identity,
                "succeeded",
                selected.translation,
                selected.stage,
                review_details=tuple(
                    {
                        "review_status": "retained",
                        "review_reason": "当前指守卫角色，保留此语境适用的译法。",
                        "review_contract_version": 1,
                        "review_reference_terms": {"Guards": "瓜尔"},
                        "review_before_text": selected.translation,
                        "detected_issues": [{"issue_type": "term_mismatch", "term": "Guards"}],
                    }.items()
                ),
            )
        }

    (closed,) = close_task_terms(
        (outcome,),
        resolve_terms=lambda *_: {"Guards": "瓜尔"},
        repair=supplement,
        stop_event=threading.Event(),
    )
    reviewed = closed.polish[original.id]
    assert seen == [{ai_entry_key(original): {"Guards": "瓜尔"}}]
    assert reviewed.original_translation == "守备"
    assert reviewed.review_before_text == reviewed.polished_translation == "卫兵"
    assert reviewed.review_status == "retained" and reviewed.accepted
    details = source_reference_review.source_review_evidence(closed)[original.identity][1]
    assert details["review_before_text"] == "卫兵"

    ledger = TaskEntryResults((task,), "latest-review")
    ledger.ingest((closed,), (original.identity,))
    snapshot = ledger.snapshot(task.key)
    assert snapshot.failure_count == 0 and snapshot.accepted_count == 1
    workbook = load_workbook(BytesIO(ExcelReportRenderer().render(snapshot).content))
    row = tuple(workbook["校对标注"].values)[1]
    assert row[3] == row[4] == "卫兵"
    assert row[5] == "保留原译：语境适用"


@pytest.mark.parametrize("review_before", [None, ""])
def test_review_before_fallback_distinguishes_missing_from_empty_text(review_before):
    entry = TranslationEntry("guard", "guard", "Guard", "原译", STAGE_TRANSLATED, "NPC_:FULL")
    result = ProofreadPipeline._proofread_result(
        entry,
        valid=True,
        translation="现译",
        review_details={
            "review_status": "revised",
            "review_reason": "根据语境修订。",
            "review_before_text": review_before,
        },
    )
    snapshot = build_polish_report_snapshot({entry.id: result}, (entry,), run_id="before-fallback")
    candidate = snapshot.candidates[0]
    assert candidate.before_text == "原译"
    assert dict(candidate.report_details)["review_before_text"] == ("原译" if review_before is None else "")
