"""Review marks have independent filters and exports stay outside the Qt event loop."""

import threading

from openpyxl import load_workbook
from PyQt6.QtCore import QThread, QTimer

from tests.contracts.translation.test_advisory_review_report import review_snapshot
from tests.ui.tools import test_ai_task_history_ui as fixtures
from transbridge.application.translation.postprocess_report import ExcelReportRenderer
from transbridge.ui.tools.ai_translator import task_result_dialog
from transbridge.ui.tools.ai_translator.task_result_dialog import TaskResultDialog
from transbridge.ui.tools.ai_translator.task_result_projection import ResultProjector

qapp = fixtures.qapp


def record():
    return {"sources": [{"key": "plugin", "label": "BSHeartland", "snapshot": review_snapshot().to_dict()}]}


def test_advisory_status_filters_do_not_replace_execution_status():
    source = record()["sources"][0]
    projector = ResultProjector()
    assert len(projector((source, "", "已接受")).visible) == 4
    pending = projector((source, "", "待复核"))
    assert len(pending.visible) == 1
    assert pending.visible[0]["accepted"] is True
    assert "AI 校对：待复核" in pending.index.details[id(pending.visible[0])]
    assert "瓜尔" in pending.index.details[id(pending.visible[0])]
    assert len(projector((source, "瓜尔", "疑似参考术语错误")).visible) == 1


def test_result_export_restores_renders_and_saves_review_evidence_in_background(qapp, tmp_path, monkeypatch):
    dialog = TaskResultDialog(record())
    entered, release = threading.Event(), threading.Event()
    threads, heartbeat = {}, []
    restore = task_result_dialog.restore_report_snapshot
    render = ExcelReportRenderer.render

    def restore_observed(data):
        threads["restore"] = QThread.currentThread()
        return restore(data)

    def render_observed(self, snapshot, **kwargs):
        threads["render"] = QThread.currentThread()
        entered.set()
        assert release.wait(3)
        return render(self, snapshot, **kwargs)

    monkeypatch.setattr(task_result_dialog, "restore_report_snapshot", restore_observed)
    monkeypatch.setattr(ExcelReportRenderer, "render", render_observed)
    target = tmp_path / "review.xlsx"
    try:
        dialog.filter.setCurrentText("待复核")
        fixtures._wait(qapp, lambda: not dialog._projecting)
        assert dialog.model.rowCount() == 1
        assert dialog.model.data(dialog.model.index(0, 4)) == "已接受"
        assert dialog.model.data(dialog.model.index(0, 6)) == "待复核"
        dialog.export(target, "xlsx")
        fixtures._wait(qapp, entered.is_set)
        QTimer.singleShot(0, lambda: heartbeat.append(True))
        fixtures._wait(qapp, lambda: bool(heartbeat))
        assert dialog.busy
        release.set()
        fixtures._wait(qapp, lambda: not dialog.busy)
        assert all(thread is not qapp.thread() for thread in threads.values())
        workbook = load_workbook(target)
        assert workbook["校对标注"].max_row == 5  # Export includes all source results, not only the filter.
        assert "已导出" in dialog.export_status.text()
    finally:
        release.set()
        fixtures._wait(qapp, lambda: not dialog.busy)
        dialog.close()
        dialog.deleteLater()
        qapp.processEvents()
